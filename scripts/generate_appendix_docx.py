"""Генерация docs/appendix.docx — компактное приложение (~15 стр.), подсветка как 123.docx."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from pygments import lex
from pygments.lexers import PythonLexer, YamlLexer
from pygments.token import Token

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "appendix.docx"

HEADER_FONT = "Times New Roman"
HEADER_SIZE = Pt(14)
CODE_FONT = "Consolas"
CODE_SIZE = Pt(10.5)

COLOR_DEFAULT = RGBColor(0x00, 0x00, 0x00)
COLOR_KEYWORD = RGBColor(0x00, 0x00, 0xFF)
COLOR_STRING = RGBColor(0xA3, 0x15, 0x15)
COLOR_COMMENT = RGBColor(0x00, 0x80, 0x00)
COLOR_NUMBER = RGBColor(0x09, 0x86, 0x58)

PYTHON_LEXER = PythonLexer()
YAML_LEXER = YamlLexer()

# Компактное приложение (~15 стр.): только ключевые фрагменты
SECTIONS: list[tuple[str, str, list[tuple[int, int]] | None]] = [
    ("docker-compose.yml", "docker-compose.yml", [(1, 10)]),
    ("app > main.py", "app/main.py", [(43, 65)]),
    ("app > models > user.py", "app/models/user.py", [(23, 39)]),
    ("app > models > request.py", "app/models/request.py", [(14, 40)]),
    ("app > models > proof.py", "app/models/proof.py", [(14, 30)]),
    ("app > schemas > request.py", "app/schemas/request.py", [(18, 38)]),
    ("app > routes > users.py", "app/routes/users.py", [(31, 38), (74, 87)]),
    ("app > routes > requests.py", "app/routes/requests.py", [(82, 93), (119, 143)]),
    ("app > routes > proofs.py", "app/routes/proofs.py", [(29, 51), (62, 76)]),
    ("app > admin_auth.py", "app/admin_auth.py", [(16, 32)]),
    ("app > bot > api_client.py", "app/bot/api_client.py", [(53, 60), (114, 121), (155, 162)]),
    ("app > bot > handlers.py", "app/bot/handlers.py", [(39, 55), (199, 215)]),
    (
        "app > bot > staff_handlers.py",
        "app/bot/staff_handlers.py",
        [(75, 90), (244, 265), (276, 290)],
    ),
]


def _token_color(ttype) -> RGBColor:
    if ttype in Token.Keyword or ttype in Token.Keyword.Constant:
        return COLOR_KEYWORD
    if ttype in Token.Name.Builtin or ttype in Token.Name.Builtin.Pseudo:
        return COLOR_KEYWORD
    if ttype in Token.Literal.String or ttype in Token.Literal.String.Doc:
        return COLOR_STRING
    if ttype in Token.Comment:
        return COLOR_COMMENT
    if ttype in Token.Literal.Number:
        return COLOR_NUMBER
    return COLOR_DEFAULT


def _highlight_by_lines(source: str, lexer) -> list[list[tuple[str, RGBColor]]]:
    if not source:
        return [[]]
    raw_lines = source.splitlines()
    line_runs: list[list[tuple[str, RGBColor]]] = [[] for _ in range(len(raw_lines) or 1)]
    pos = 0
    for ttype, value in lex(source, lexer):
        if not value:
            continue
        color = _token_color(ttype)
        idx = source.find(value, pos)
        if idx < 0:
            idx = pos
        pos = idx + len(value)
        start_line = source.count("\n", 0, idx)
        if start_line < len(line_runs):
            line_runs[start_line].append((value, color))
    merged: list[list[tuple[str, RGBColor]]] = []
    for i, raw in enumerate(raw_lines):
        runs = line_runs[i] if i < len(line_runs) else []
        if not runs:
            merged.append([(raw, COLOR_DEFAULT)] if raw else [])
        elif "".join(t for t, _ in runs) == raw:
            merged.append(runs)
        else:
            merged.append(_fill_gaps(raw, runs))
    return merged


def _fill_gaps(raw: str, runs: list[tuple[str, RGBColor]]) -> list[tuple[str, RGBColor]]:
    if not raw:
        return []
    out: list[tuple[str, RGBColor]] = []
    idx = 0
    for text, color in runs:
        start = raw.find(text, idx)
        if start == -1:
            continue
        if start > idx:
            out.append((raw[idx:start], COLOR_DEFAULT))
        out.append((text, color))
        idx = start + len(text)
    if idx < len(raw):
        out.append((raw[idx:], COLOR_DEFAULT))
    return out or [(raw, COLOR_DEFAULT)]


def _read_excerpt(rel: str, ranges: list[tuple[int, int]] | None) -> str:
    lines = (ROOT / rel).read_text(encoding="utf-8").splitlines()
    if ranges is None:
        return "\n".join(lines)
    parts: list[str] = []
    for i, (start, end) in enumerate(ranges):
        if i:
            parts.append("")
        parts.extend(lines[start - 1 : end])
    return "\n".join(parts)


def add_title(doc: Document, text: str, *, bold: bool = True) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.name = HEADER_FONT
    run.font.size = HEADER_SIZE


def add_section(doc: Document, header: str, rel: str, ranges: list[tuple[int, int]] | None) -> None:
    source = _read_excerpt(rel, ranges)
    add_title(doc, header, bold=False)
    lexer = YAML_LEXER if rel.endswith((".yml", ".yaml")) else PYTHON_LEXER
    for runs in _highlight_by_lines(source, lexer):
        p = doc.add_paragraph()
        if not runs:
            continue
        for text, color in runs:
            run = p.add_run(text)
            run.font.name = CODE_FONT
            run.font.size = CODE_SIZE
            run.font.color.rgb = color
    doc.add_paragraph()


def main() -> None:
    doc = Document()
    doc.add_paragraph()
    add_title(doc, "ПРИЛОЖЕНИЕ")

    total_lines = 0
    for header, rel, ranges in SECTIONS:
        add_section(doc, header, rel, ranges)
        excerpt = _read_excerpt(rel, ranges)
        total_lines += len(excerpt.splitlines())

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Saved: {OUT} (~{total_lines} lines of code)")


if __name__ == "__main__":
    main()
